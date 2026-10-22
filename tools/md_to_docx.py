#!/usr/bin/env python3
"""把项目 Markdown 技术文档渲染成排版规范的 DOCX。

只支持本项目实际用到的 Markdown 子集：
  标题(#..####)、段落、无序列表(-)、有序列表(1.)、引用(>)、
  围栏代码块(```)、GitHub 管道表格、分隔线(---)、行内 **粗体** 与 `代码`。
"""
import os
import re
import sys
import math

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT_CN = "宋体"
BODY_FONT_EN = "Calibri"
HEAD_FONT_CN = "黑体"
CODE_FONT = "Consolas"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

# OOXML 对 pPr / tcPr 的子元素顺序有严格要求，顺序错了 Word 可能提示"文件已损坏"。
# 这里在所有修改完成后按 schema 顺序重排一次，避免手工 append 造成乱序。
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
    "cellDel", "cellMerge", "tcPrChange",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
    "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]


def reorder_children(el, order):
    children = list(el)
    indexed = []
    for child in children:
        tag = child.tag.split("}")[-1]
        pos = order.index(tag) if tag in order else len(order)
        indexed.append((pos, child))
        el.remove(child)
    for _, child in sorted(indexed, key=lambda x: x[0]):
        el.append(child)


def normalize_order(doc):
    body = doc.element.body
    for ppr in body.iter(qn("w:pPr")):
        reorder_children(ppr, PPR_ORDER)
    for tcpr in body.iter(qn("w:tcPr")):
        reorder_children(tcpr, TCPR_ORDER)
    for tblpr in body.iter(qn("w:tblPr")):
        reorder_children(tblpr, TBLPR_ORDER)
    for trpr in body.iter(qn("w:trPr")):
        reorder_children(trpr, TRPR_ORDER)


def set_run_font(run, cn=BODY_FONT_CN, en=BODY_FONT_EN, size=None, bold=None, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(element, fill):
    ppr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_inline(paragraph, text, size=10.5):
    """解析 **粗体** 与 `行内代码`。"""
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            set_run_font(run, size=size, bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, en=CODE_FONT, cn=CODE_FONT, size=size - 0.5)
            run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x1D)
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, size=size)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT_EN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CN)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)

    sizes = {"Heading 1": 16, "Heading 2": 13, "Heading 3": 11.5}
    for name, size in sizes.items():
        st = doc.styles[name]
        st.font.name = BODY_FONT_EN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT_CN)
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def add_page_number_footer(doc, text):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text + "    第 ")
    set_run_font(run, size=8.5, color=(0x88, 0x88, 0x88))

    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

    run = p.add_run(" 页")
    set_run_font(run, size=8.5, color=(0x88, 0x88, 0x88))


def add_code_block(doc, lines):
    if not lines:
        return
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.left_indent = Cm(0.3)
    shade(p._p, "F5F5F5")
    # 左边框，视觉上区分代码块
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "C8C8C8")
    borders.append(left)
    ppr.append(borders)

    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, en=CODE_FONT, cn=CODE_FONT, size=8.5)
        if idx != len(lines) - 1:
            run.add_break()


def set_table_width_pct(table, pct):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(pct))
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)   # 顺序由 normalize_order 统一重排


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def display_len(text):
    """按显示宽度估算：CJK 记 2，ASCII 记 1。"""
    n = 0
    for ch in text:
        n += 2 if ord(ch) > 0x2E7F else 1
    return n


def compute_widths(header, body, total_cm=15.2):
    """按内容长度开方加权分配列宽：长文本列宽，短列窄，避免大段空白。"""
    cols = len(header)
    maxlen = [max(4, display_len(h)) for h in header]
    for row in body:
        for i in range(cols):
            maxlen[i] = max(maxlen[i], display_len(row[i]))
    weights = [math.sqrt(v) for v in maxlen]
    s = sum(weights)
    widths = [max(1.5, total_cm * w / s) for w in weights]
    s2 = sum(widths)
    return [w * total_cm / s2 for w in widths]


def add_table(doc, rows):
    if len(rows) < 2:
        return
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = []
    for r in rows[2:]:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if len(cells) < len(header):
            cells += [""] * (len(header) - len(cells))
        body.append(cells[: len(header)])

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_width_pct(table, 5000)   # 表格占满正文宽度

    hdr = table.rows[0]
    no_split(hdr)                      # 表头行也不允许跨页拆开
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, size=9)
        for run in p.runs:
            run.font.bold = True
        shade(cell._tc.get_or_add_tcPr(), "EDEDED")

    # 表头跨页重复
    trpr = hdr._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)

    for row in body:
        new_row = table.add_row()
        no_split(new_row)              # 防止一行文字被分页切成两半
        cells = new_row.cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, text, size=9)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.6

    # 固定列宽：autofit 会把各列拉成均分，导致短列大量留白、长列频繁折行
    widths = compute_widths(header, body)
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for i, column in enumerate(table.columns):
        if i < len(widths):
            column.width = Cm(widths[i])      # 同步 tblGrid，LibreOffice 依此排版
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Cm(widths[i])


def add_quote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    add_inline(p, text, size=9.5)
    for run in p.runs:
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D0D0D0")
    borders.append(bottom)
    ppr.append(borders)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style_document(doc)
    core = doc.core_properties
    core.title = "高性能 Reactor 网络库技术设计说明"
    core.subject = "C++11 / epoll / 主从 Reactor / 线程池 / 二进制协议"
    core.author = "MiniNet"

    i = 0
    total = len(lines)
    first_heading_done = False

    while i < total:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 图片：![caption](relative/path.png)
        m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", stripped)
        if m:
            caption, rel = m.group(1), m.group(2)
            img_path = os.path.join(os.path.dirname(md_path), rel.replace("/", os.sep))
            if os.path.isfile(img_path):
                doc.add_picture(img_path, width=Cm(15.2))
                pic_par = doc.paragraphs[-1]
                pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic_par.paragraph_format.space_before = Pt(6)
                pic_par.paragraph_format.space_after = Pt(2)
                pic_par.paragraph_format.keep_with_next = True

                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(10)
                run = cap.add_run(caption)
                set_run_font(run, size=9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # 代码块
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < total and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code)
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < total and re.match(r"^\|[\s:\-|]+\|?$", lines[i + 1].strip()):
            block = []
            while i < total and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_heading_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(18)
                run = p.add_run(text)
                set_run_font(run, cn=HEAD_FONT_CN, size=22, bold=True)
                run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
                first_heading_done = True
            else:
                # 文档标题单独处理，因此 ## 作为一级标题，保持导航层级干净
                doc.add_heading(text, level=max(1, min(level, 4) - 1))
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", stripped):
            add_rule(doc)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            add_quote(doc, stripped.lstrip(">").strip())
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        # 有序列表：手工编号，避免 Word 在多段列表间连续计数
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, f"{m.group(1)}. {m.group(2)}")
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_inline(p, stripped)
        i += 1

    add_page_number_footer(doc, "MiniNet 技术设计说明")
    normalize_order(doc)
    doc.save(docx_path)
    return doc


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md = os.path.join(root, "docs", "技术设计说明.md")
    out = os.path.join(root, "docs", "高性能Reactor网络库技术设计说明.docx")
    doc = convert(md, out)
    print(f"saved: {out}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
